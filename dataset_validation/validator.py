"""
validator.py
============
Dataset validation engine with scoring (1 or 0 per column), row total scores,
date tolerance pass rules, and baseline regression comparison (+1 / -1 / 0 per feature).

Rules:
  - Column score is 1 (if match / fuzz >= 95%) or 0.
  - Rule 1: If is_certificate is False in golden and model says False -> 1 in every column.
  - Rule 2: If is_iso_iatf is False in golden and model says False -> 1 in every remaining column.
  - Row pass/fail: PASS if all 10 columns are correct, OR if all correct except exactly 1 date column.
  - Baseline comparison: Compares feature scores with analytics/baseline.json.
    If total delta > 0, saves to analytics/new_baseline.json.
"""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Dict, Any, List, Optional, Tuple

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .rules import (
    compare_boolean,
    compare_text_rapidfuzz,
    compare_site_iso_type,
    compare_dates,
    normalize_text,
    normalize_bool,
)
from verifier import CertificateVerifier

EVALUATED_COLUMNS = [
    "is_certificate",
    "is_iso_iatf",
    "cert_type",
    "site_name",
    "address",
    "scope",
    "iso_type",
    "issued_date",
    "expiration_date",
    "site_iso_type",
]


class ColumnEvaluation:
    def __init__(
        self,
        column_name: str,
        expected: Any,
        predicted: Any,
        passed: bool,
        score: int,  # 1 or 0
        similarity_percent: float = 100.0,
        note: str = "",
    ):
        self.column_name = column_name
        self.expected = expected
        self.predicted = predicted
        self.passed = passed
        self.score = 1 if score else 0
        self.similarity_percent = round(similarity_percent, 2)
        self.note = note

    def to_dict(self) -> Dict[str, Any]:
        return {
            "expected": self.expected,
            "predicted": self.predicted,
            "passed": self.passed,
            "score": self.score,
            "similarity_percent": self.similarity_percent,
            "note": self.note,
        }


class RowEvaluation:
    def __init__(self, url: str, index: int):
        self.url = url
        self.index = index
        self.overall_passed = True
        self.date_tolerance_applied = False
        self.total_score = 0
        self.avg_score = 0.0
        self.columns: Dict[str, ColumnEvaluation] = {}
        self.verifier_status: Optional[str] = None
        self.confidence_score: Optional[int] = None
        self.verification_reviews: Optional[str] = None
        self.is_false_pass: bool = False

    def add_column_result(self, col_eval: ColumnEvaluation):
        self.columns[col_eval.column_name] = col_eval

    def finalize(self):
        """
        Computes row total score, avg score, overall pass/fail status,
        and false pass detection (verifier Approved vs validator Failed).
        Date tolerance: If all columns are correct except exactly 1 date column
        (either issued_date or expiration_date), the row PASSES.
        """
        self.total_score = sum(c.score for c in self.columns.values())
        self.avg_score = round(self.total_score / float(len(EVALUATED_COLUMNS)), 4)

        if self.total_score == len(EVALUATED_COLUMNS):
            self.overall_passed = True
            self.date_tolerance_applied = False
        elif self.total_score == len(EVALUATED_COLUMNS) - 1:
            # Check which column failed
            failed_cols = [c_name for c_name, c in self.columns.items() if c.score == 0]
            if len(failed_cols) == 1 and failed_cols[0] in ("issued_date", "expiration_date"):
                self.overall_passed = True
                self.date_tolerance_applied = True
            else:
                self.overall_passed = False
                self.date_tolerance_applied = False
        else:
            self.overall_passed = False
            self.date_tolerance_applied = False

        # False Pass: verifier says Approved, but validator evaluated overall_passed as False
        self.is_false_pass = bool(self.verifier_status == "Approved" and not self.overall_passed)

    def to_dict(self) -> Dict[str, Any]:
        return {
            "index": self.index,
            "url": self.url,
            "overall_passed": self.overall_passed,
            "date_tolerance_applied": self.date_tolerance_applied,
            "verifier_status": self.verifier_status,
            "confidence_score": self.confidence_score,
            "verification_reviews": self.verification_reviews,
            "is_false_pass": self.is_false_pass,
            "total_score": self.total_score,
            "avg_score": self.avg_score,
            "column_scores": {k: v.score for k, v in self.columns.items()},
            "columns": {k: v.to_dict() for k, v in self.columns.items()},
        }


class DatasetValidator:
    """
    Validates predicted extractor DataFrame against expected golden dataset DataFrame.
    Scores each column (1 or 0), calculates row total and averages, applies date tolerance,
    and supports baseline regression tracking.
    """

    def __init__(self, golden_path: str):
        self.golden_path = golden_path
        self.golden_df = pd.read_excel(golden_path)

    def _normalize_url_key(self, url: Any) -> str:
        s = str(url or "").strip().rstrip("/")
        return s.lower()

    def validate(self, predicted_df: pd.DataFrame) -> Dict[str, Any]:
        """
        Runs rule-based validation on predicted_df against golden_df.
        Returns full validation summary and row evaluations.
        """
        pred_map: Dict[str, Dict[str, Any]] = {}
        for idx, row in predicted_df.iterrows():
            url = row.get("url")
            if url and str(url).strip():
                key = self._normalize_url_key(url)
                pred_map[key] = row.to_dict()

        row_evaluations: List[RowEvaluation] = []

        for idx, expected_row in self.golden_df.iterrows():
            url = str(expected_row.get("url") or "").strip()
            key = self._normalize_url_key(url)

            if key in pred_map:
                predicted_row = pred_map[key]
            elif idx < len(predicted_df):
                predicted_row = predicted_df.iloc[idx].to_dict()
            else:
                predicted_row = {}

            row_eval = self._evaluate_single_row(idx + 1, url, expected_row, predicted_row)
            row_eval.finalize()
            row_evaluations.append(row_eval)

        summary = self._build_summary(row_evaluations)
        return {
            "summary": summary,
            "rows": row_evaluations,
        }

    def _evaluate_single_row(
        self,
        row_idx: int,
        url: str,
        expected: Any,
        predicted: Any,
    ) -> RowEvaluation:
        row_eval = RowEvaluation(url=url, index=row_idx)

        # ── Step 0: Extract or evaluate verification metadata ────────────────
        pred_dict = (
            predicted if isinstance(predicted, dict)
            else (predicted.to_dict() if hasattr(predicted, "to_dict") else {})
        )
        v_status = pred_dict.get("verification_status")
        c_score = pred_dict.get("confidence_score")
        v_reviews = pred_dict.get("verification_reviews")

        if not v_status and pred_dict:
            verif_res = CertificateVerifier().verify_row(pred_dict)
            v_status = verif_res.status
            c_score = verif_res.score
            v_reviews = verif_res.reviews

        row_eval.verifier_status = v_status
        row_eval.confidence_score = c_score
        row_eval.verification_reviews = v_reviews

        # ── Step 1: Rule 1 - is_certificate ──────────────────────────────────
        exp_is_cert = normalize_bool(expected.get("is_certificate"))
        pred_is_cert = normalize_bool(predicted.get("is_certificate"))
        cert_match, cert_msg = compare_boolean(exp_is_cert, pred_is_cert)
        cert_score = 1 if cert_match else 0

        row_eval.add_column_result(
            ColumnEvaluation(
                column_name="is_certificate",
                expected=exp_is_cert,
                predicted=pred_is_cert,
                passed=cert_match,
                score=cert_score,
                similarity_percent=100.0 if cert_match else 0.0,
                note=cert_msg,
            )
        )

        # RULE 1: If is_certificate is False in the golden dataset:
        # If model correctly predicted False -> get 1 in every column.
        if exp_is_cert is False:
            if pred_is_cert is False:
                for col in EVALUATED_COLUMNS[1:]:
                    row_eval.add_column_result(
                        ColumnEvaluation(
                            column_name=col,
                            expected=expected.get(col),
                            predicted=predicted.get(col),
                            passed=True,
                            score=1,
                            similarity_percent=100.0,
                            note="Rule 1 awarded (is_certificate is False)",
                        )
                    )
                return row_eval
            else:
                # Model incorrectly predicted True for is_certificate (score 0 already)
                # Remaining columns expected to be empty.
                for col in EVALUATED_COLUMNS[1:]:
                    val_pred = predicted.get(col)
                    is_empty = val_pred is None or pd.isna(val_pred) or str(val_pred).strip().lower() in ("", "none", "nan")
                    row_eval.add_column_result(
                        ColumnEvaluation(
                            column_name=col,
                            expected=expected.get(col),
                            predicted=val_pred,
                            passed=is_empty,
                            score=1 if is_empty else 0,
                            similarity_percent=100.0 if is_empty else 0.0,
                            note="Expected empty (is_certificate False)" if not is_empty else "Correctly empty",
                        )
                    )
                return row_eval

        # ── Step 2: Rule 2 - is_iso_iatf ─────────────────────────────────────
        exp_is_iso = normalize_bool(expected.get("is_iso_iatf"))
        pred_is_iso = normalize_bool(predicted.get("is_iso_iatf"))
        iso_match, iso_msg = compare_boolean(exp_is_iso, pred_is_iso)
        iso_score = 1 if iso_match else 0

        row_eval.add_column_result(
            ColumnEvaluation(
                column_name="is_iso_iatf",
                expected=exp_is_iso,
                predicted=pred_is_iso,
                passed=iso_match,
                score=iso_score,
                similarity_percent=100.0 if iso_match else 0.0,
                note=iso_msg,
            )
        )

        # RULE 2: If is_iso_iatf is False in the golden dataset:
        # If model correctly predicted False -> get 1 in every remaining column.
        if exp_is_iso is False:
            if pred_is_iso is False:
                for col in EVALUATED_COLUMNS[2:]:
                    row_eval.add_column_result(
                        ColumnEvaluation(
                            column_name=col,
                            expected=expected.get(col),
                            predicted=predicted.get(col),
                            passed=True,
                            score=1,
                            similarity_percent=100.0,
                            note="Rule 2 awarded (is_iso_iatf is False)",
                        )
                    )
                return row_eval
            else:
                # Model incorrectly predicted True for is_iso_iatf
                for col in EVALUATED_COLUMNS[2:]:
                    val_pred = predicted.get(col)
                    is_empty = val_pred is None or pd.isna(val_pred) or str(val_pred).strip().lower() in ("", "none", "nan")
                    row_eval.add_column_result(
                        ColumnEvaluation(
                            column_name=col,
                            expected=expected.get(col),
                            predicted=val_pred,
                            passed=is_empty,
                            score=1 if is_empty else 0,
                            similarity_percent=100.0 if is_empty else 0.0,
                            note="Expected empty (is_iso_iatf False)" if not is_empty else "Correctly empty",
                        )
                    )
                return row_eval

        # ── Step 3: Rules 3, 4, 5 for remaining columns ─────────────────────
        # Text fields checked with RapidFuzz >= 95% -> 1, else 0
        for col in ["cert_type", "site_name", "address", "scope", "iso_type"]:
            exp_val = expected.get(col)
            pred_val = predicted.get(col)
            is_match, score_pct, note = compare_text_rapidfuzz(exp_val, pred_val, threshold=95.0)
            row_eval.add_column_result(
                ColumnEvaluation(
                    column_name=col,
                    expected=exp_val,
                    predicted=pred_val,
                    passed=is_match,
                    score=1 if is_match else 0,
                    similarity_percent=score_pct,
                    note=note,
                )
            )

        # Dates checked with date normalization
        for date_col in ["issued_date", "expiration_date"]:
            exp_d = expected.get(date_col)
            pred_d = predicted.get(date_col)
            is_match, score_pct, note = compare_dates(exp_d, pred_d)
            row_eval.add_column_result(
                ColumnEvaluation(
                    column_name=date_col,
                    expected=exp_d,
                    predicted=pred_d,
                    passed=is_match,
                    score=1 if is_match else 0,
                    similarity_percent=score_pct,
                    note=note,
                )
            )

        # Rule 4: site_iso_type checked with exact match
        exp_iso_type = expected.get("site_iso_type")
        pred_iso_type = predicted.get("site_iso_type")
        is_match, score_pct, note = compare_site_iso_type(exp_iso_type, pred_iso_type)
        row_eval.add_column_result(
            ColumnEvaluation(
                column_name="site_iso_type",
                expected=exp_iso_type,
                predicted=pred_iso_type,
                passed=is_match,
                score=1 if is_match else 0,
                similarity_percent=score_pct,
                note=note,
            )
        )

        return row_eval

    def _build_summary(self, row_evaluations: List[RowEvaluation]) -> Dict[str, Any]:
        total_rows = len(row_evaluations)
        passed_rows = sum(1 for r in row_evaluations if r.overall_passed)
        failed_rows = total_rows - passed_rows
        pass_rate = round((passed_rows / total_rows * 100), 2) if total_rows > 0 else 0.0

        # False Pass calculation:
        # Rows where verifier approved the certificate (confidence >= 90) but validator evaluated overall_passed as False
        false_pass_rows = [r for r in row_evaluations if r.is_false_pass]
        false_pass_cases = len(false_pass_rows)
        false_pass_rate_total = round((false_pass_cases / total_rows * 100), 2) if total_rows > 0 else 0.0
        false_pass_rate_failed = round((false_pass_cases / failed_rows * 100), 2) if failed_rows > 0 else 0.0

        total_row_scores = [r.total_score for r in row_evaluations]
        avg_row_scores = [r.avg_score for r in row_evaluations]

        avg_total_score = round(sum(total_row_scores) / float(total_rows), 4) if total_rows > 0 else 0.0
        overall_avg_row_score = round(sum(avg_row_scores) / float(total_rows), 4) if total_rows > 0 else 0.0

        column_stats: Dict[str, Any] = {}
        feature_scores: Dict[str, float] = {}

        for col in EVALUATED_COLUMNS:
            scores = [r.columns[col].score for r in row_evaluations if col in r.columns]
            sim_scores = [r.columns[col].similarity_percent for r in row_evaluations if col in r.columns]

            total_pts = sum(scores)
            avg_score = round(total_pts / float(total_rows), 4) if total_rows > 0 else 0.0
            avg_sim = round(sum(sim_scores) / float(len(sim_scores)), 2) if sim_scores else 0.0

            feature_scores[f"avg_{col}_score"] = avg_score
            column_stats[col] = {
                "total_points": total_pts,
                "max_possible_points": total_rows,
                "avg_score": avg_score,
                "accuracy_percent": round(avg_score * 100, 2),
                "avg_similarity_percent": avg_sim,
            }

        return {
            "validation_timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "golden_file": self.golden_path,
            "total_cases": total_rows,
            "passed_cases": passed_rows,
            "failed_cases": failed_rows,
            "pass_rate_percent": pass_rate,
            "false_pass_cases": false_pass_cases,
            "false_pass_rate_total_percent": false_pass_rate_total,
            "false_pass_rate_failed_percent": false_pass_rate_failed,
            "false_pass_urls": [
                {
                    "index": r.index,
                    "url": r.url,
                    "confidence_score": r.confidence_score,
                    "total_score": r.total_score,
                    "verification_reviews": r.verification_reviews,
                }
                for r in false_pass_rows
            ],
            "average_total_score": avg_total_score,
            "average_row_score": overall_avg_row_score,
            "feature_scores": feature_scores,
            "column_breakdown": column_stats,
        }

    def compare_with_baseline(
        self,
        current_summary: Dict[str, Any],
        analytics_dir: str = "analytics",
        baseline_filename: str = "baseline.json",
        new_baseline_filename: str = "new_baseline.json",
    ) -> Dict[str, Any]:
        """
        Compares current run metrics against analytics/baseline.json.
        Rules:
          - If a feature has improved -> +1
          - If a feature has decreased -> -1
          - If the same -> +0
          - False Pass Rate:
            If decreased (fewer false passes) -> +1 (Improved)
            If increased (more false passes)   -> -1 (Decreased / Regressed)
            If the same                        -> +0 (Same)
        If total improvement delta is positive (> 0), saves current run to analytics/new_baseline.json.
        """
        dir_path = Path(analytics_dir)
        dir_path.mkdir(parents=True, exist_ok=True)
        baseline_path = dir_path / baseline_filename
        new_baseline_path = dir_path / new_baseline_filename

        # If baseline does not exist, save current run as initial baseline
        if not baseline_path.exists():
            with open(baseline_path, "w", encoding="utf-8") as f:
                json.dump(current_summary, f, indent=2, ensure_ascii=False)
            return {
                "has_baseline": False,
                "baseline_saved": True,
                "baseline_path": str(baseline_path.resolve()),
                "total_delta": 0,
                "feature_comparisons": {},
                "saved_new_baseline": False,
            }

        # Baseline exists -> Load and compare
        with open(baseline_path, "r", encoding="utf-8") as f:
            baseline_data = json.load(f)

        base_features = baseline_data.get("feature_scores", {})
        curr_features = current_summary.get("feature_scores", {})

        feature_comparisons: Dict[str, Any] = {}
        total_delta = 0

        for feat_name, curr_val in curr_features.items():
            base_val = base_features.get(feat_name, 0.0)
            diff = round(curr_val - base_val, 4)

            if diff > 0:
                delta = 1
                status = "+1 (Improved)"
            elif diff < 0:
                delta = -1
                status = "-1 (Decreased)"
            else:
                delta = 0
                status = "+0 (Same)"

            total_delta += delta
            feature_comparisons[feat_name] = {
                "baseline": base_val,
                "new": curr_val,
                "difference": diff,
                "delta": delta,
                "status": status,
            }

        # Compare False Pass Rate as a feature
        # If false_pass_rate decreases -> +1 (Improvement)
        # If false_pass_rate increases -> -1 (Regression)
        # If false_pass_rate is same   -> +0 (Same)
        base_fp_rate = float(baseline_data.get("false_pass_rate_total_percent", 0.0))
        curr_fp_rate = float(current_summary.get("false_pass_rate_total_percent", 0.0))
        diff_fp = round(curr_fp_rate - base_fp_rate, 2)

        if diff_fp < 0:
            fp_delta = 1
            fp_status = "+1 (Improved - fewer false passes)"
        elif diff_fp > 0:
            fp_delta = -1
            fp_status = "-1 (Decreased - more false passes)"
        else:
            fp_delta = 0
            fp_status = "+0 (Same)"

        total_delta += fp_delta
        feature_comparisons["false_pass_rate"] = {
            "baseline": base_fp_rate,
            "new": curr_fp_rate,
            "difference": diff_fp,
            "delta": fp_delta,
            "status": fp_status,
        }

        saved_new_baseline = False
        if total_delta > 0:
            with open(new_baseline_path, "w", encoding="utf-8") as f:
                json.dump(current_summary, f, indent=2, ensure_ascii=False)
            saved_new_baseline = True

        return {
            "has_baseline": True,
            "baseline_path": str(baseline_path.resolve()),
            "total_delta": total_delta,
            "is_positive": total_delta > 0,
            "feature_comparisons": feature_comparisons,
            "saved_new_baseline": saved_new_baseline,
            "new_baseline_path": str(new_baseline_path.resolve()) if saved_new_baseline else None,
        }

    def save_reports(
        self,
        results: Dict[str, Any],
        output_excel: str = "dataset_validation/validation_report.xlsx",
        output_json: str = "dataset_validation/validation_summary.json",
    ) -> Tuple[str, str]:
        """
        Saves the validation report as a styled Excel file and a JSON summary.
        """
        out_excel_path = Path(output_excel)
        out_excel_path.parent.mkdir(parents=True, exist_ok=True)
        out_json_path = Path(output_json)
        out_json_path.parent.mkdir(parents=True, exist_ok=True)

        summary = results["summary"]
        rows: List[RowEvaluation] = results["rows"]

        # ── 1. Save JSON ──────────────────────────────────────────────────────
        json_data = {
            "summary": summary,
            "rows": [r.to_dict() for r in rows],
        }
        with open(out_json_path, "w", encoding="utf-8") as f:
            json.dump(json_data, f, indent=2, ensure_ascii=False, default=str)

        # ── 2. Save Styled Excel Report ───────────────────────────────────────
        wb = Workbook()
        ws_summary = wb.active
        ws_summary.title = "Validation Summary"

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=14, bold=True, color="1F497D")
        bold_font = Font(name="Calibri", size=11, bold=True)
        normal_font = Font(name="Calibri", size=10)

        blue_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        pass_fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
        fail_fill = PatternFill(start_color="FCE5CD", end_color="FCE5CD", fill_type="solid")
        date_tol_fill = PatternFill(start_color="D0E0E3", end_color="D0E0E3", fill_type="solid")  # soft cyan/blue for date tolerance pass

        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        # Write Summary Header
        ws_summary.cell(row=1, column=1, value="Dataset Validation Summary Report").font = title_font
        ws_summary.cell(row=2, column=1, value=f"Evaluated At: {summary['validation_timestamp']}").font = normal_font
        ws_summary.cell(row=3, column=1, value=f"Golden Dataset: {summary['golden_file']}").font = normal_font

        ws_summary.cell(row=5, column=1, value="Metric").font = header_font
        ws_summary.cell(row=5, column=1).fill = blue_fill
        ws_summary.cell(row=5, column=2, value="Value").font = header_font
        ws_summary.cell(row=5, column=2).fill = blue_fill

        summary_metrics = [
            ("Total Cases Evaluated", summary["total_cases"]),
            ("Passed Cases", summary["passed_cases"]),
            ("Failed Cases", summary["failed_cases"]),
            ("Pass Rate (%)", f"{summary['pass_rate_percent']}%"),
            ("False Pass Cases (Verifier Approved vs Validator Failed)", summary.get("false_pass_cases", 0)),
            ("False Pass Rate (of Total Cases)", f"{summary.get('false_pass_rate_total_percent', 0.0)}%"),
            ("False Pass Rate (of Failed Cases)", f"{summary.get('false_pass_rate_failed_percent', 0.0)}%"),
            ("Average Total Score (out of 10)", f"{summary['average_total_score']} / 10"),
            ("Average Row Score (0.0 - 1.0)", f"{summary['average_row_score']}"),
        ]

        curr_row = 6
        for m_name, m_val in summary_metrics:
            c1 = ws_summary.cell(row=curr_row, column=1, value=m_name)
            c2 = ws_summary.cell(row=curr_row, column=2, value=m_val)
            c1.font = bold_font
            c2.font = bold_font
            c1.border = thin_border
            c2.border = thin_border
            curr_row += 1

        # Per-column breakdown table
        curr_row += 2
        ws_summary.cell(row=curr_row, column=1, value="Column Name").font = header_font
        ws_summary.cell(row=curr_row, column=1).fill = blue_fill
        ws_summary.cell(row=curr_row, column=2, value="Feature Name").font = header_font
        ws_summary.cell(row=curr_row, column=2).fill = blue_fill
        ws_summary.cell(row=curr_row, column=3, value="Total Score Points").font = header_font
        ws_summary.cell(row=curr_row, column=3).fill = blue_fill
        ws_summary.cell(row=curr_row, column=4, value="Max Points").font = header_font
        ws_summary.cell(row=curr_row, column=4).fill = blue_fill
        ws_summary.cell(row=curr_row, column=5, value="Avg Feature Score").font = header_font
        ws_summary.cell(row=curr_row, column=5).fill = blue_fill
        ws_summary.cell(row=curr_row, column=6, value="Accuracy (%)").font = header_font
        ws_summary.cell(row=curr_row, column=6).fill = blue_fill
        ws_summary.cell(row=curr_row, column=7, value="Avg Similarity (%)").font = header_font
        ws_summary.cell(row=curr_row, column=7).fill = blue_fill

        curr_row += 1
        for col_name, stats in summary["column_breakdown"].items():
            ws_summary.cell(row=curr_row, column=1, value=col_name).font = normal_font
            ws_summary.cell(row=curr_row, column=2, value=f"avg_{col_name}_score").font = normal_font
            ws_summary.cell(row=curr_row, column=3, value=stats["total_points"]).font = normal_font
            ws_summary.cell(row=curr_row, column=4, value=stats["max_possible_points"]).font = normal_font

            score_cell = ws_summary.cell(row=curr_row, column=5, value=stats["avg_score"])
            score_cell.font = bold_font

            acc_cell = ws_summary.cell(row=curr_row, column=6, value=f"{stats['accuracy_percent']}%")
            acc_cell.font = bold_font
            acc_cell.fill = pass_fill if stats["accuracy_percent"] >= 90.0 else fail_fill

            sim_cell = ws_summary.cell(row=curr_row, column=7, value=f"{stats['avg_similarity_percent']}%")
            sim_cell.font = normal_font

            for c in range(1, 8):
                ws_summary.cell(row=curr_row, column=c).border = thin_border
            curr_row += 1

        # ── Sheet 2: Row Details ──────────────────────────────────────────────
        ws_details = wb.create_sheet(title="Row Details")
        detail_headers = [
            "Row #",
            "URL",
            "Overall Result",
            "Total Score (out of 10)",
            "Avg Score (0-1)",
            "Date Tolerance Applied?",
            "Verification Status",
            "Confidence Score",
            "False Pass? (YES/NO)",
        ]
        for col in EVALUATED_COLUMNS:
            detail_headers.extend([
                f"{col} (Expected)",
                f"{col} (Predicted)",
                f"{col} (Score: 1/0)",
                f"{col} (Sim %)",
            ])

        for col_idx, h in enumerate(detail_headers, start=1):
            cell = ws_details.cell(row=1, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = blue_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        fill_review = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        fill_false_pass = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        font_false_pass = Font(name="Calibri", size=10, bold=True, color="9C0006")

        d_row = 2
        for r in rows:
            ws_details.cell(row=d_row, column=1, value=r.index).font = normal_font
            ws_details.cell(row=d_row, column=2, value=r.url).font = normal_font

            status_str = "PASS" if r.overall_passed else "FAIL"
            res_cell = ws_details.cell(row=d_row, column=3, value=status_str)
            res_cell.font = bold_font
            res_cell.fill = (date_tol_fill if r.date_tolerance_applied else (pass_fill if r.overall_passed else fail_fill))
            res_cell.alignment = Alignment(horizontal="center")

            ws_details.cell(row=d_row, column=4, value=r.total_score).font = bold_font
            ws_details.cell(row=d_row, column=5, value=r.avg_score).font = bold_font
            ws_details.cell(row=d_row, column=6, value="YES" if r.date_tolerance_applied else "NO").font = normal_font

            # Verification Status (column 7)
            v_stat = r.verifier_status or "N/A"
            v_cell = ws_details.cell(row=d_row, column=7, value=v_stat)
            v_cell.font = bold_font
            v_cell.alignment = Alignment(horizontal="center")
            if v_stat == "Approved":
                v_cell.fill = pass_fill
            elif v_stat == "Need Review":
                v_cell.fill = fill_review
            elif v_stat == "Reject":
                v_cell.fill = fail_fill

            # Confidence Score (column 8)
            sc_val = r.confidence_score if r.confidence_score is not None else "N/A"
            sc_cell = ws_details.cell(row=d_row, column=8, value=sc_val)
            sc_cell.font = bold_font
            sc_cell.alignment = Alignment(horizontal="center")

            # False Pass? (column 9)
            fp_cell = ws_details.cell(row=d_row, column=9, value="YES (False Pass)" if r.is_false_pass else "NO")
            fp_cell.alignment = Alignment(horizontal="center")
            if r.is_false_pass:
                fp_cell.fill = fill_false_pass
                fp_cell.font = font_false_pass
            else:
                fp_cell.font = normal_font

            c_idx = 10
            for col in EVALUATED_COLUMNS:
                col_eval = r.columns.get(col)
                if col_eval:
                    ws_details.cell(row=d_row, column=c_idx, value=str(col_eval.expected if col_eval.expected is not None else "")).font = normal_font
                    ws_details.cell(row=d_row, column=c_idx + 1, value=str(col_eval.predicted if col_eval.predicted is not None else "")).font = normal_font

                    score_cell = ws_details.cell(row=d_row, column=c_idx + 2, value=col_eval.score)
                    score_cell.font = bold_font
                    score_cell.fill = pass_fill if col_eval.score == 1 else fail_fill
                    score_cell.alignment = Alignment(horizontal="center")

                    sim_cell = ws_details.cell(row=d_row, column=c_idx + 3, value=f"{col_eval.similarity_percent}%")
                    sim_cell.font = normal_font
                c_idx += 4

            for c in range(1, len(detail_headers) + 1):
                ws_details.cell(row=d_row, column=c).border = thin_border
            d_row += 1

        for ws in (ws_summary, ws_details):
            for col in ws.columns:
                col_letter = get_column_letter(col[0].column)
                max_len = max((len(str(cell.value or "")) for cell in col), default=10)
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

        wb.save(out_excel_path)
        return str(out_excel_path.resolve()), str(out_json_path.resolve())
