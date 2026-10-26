"""
main.py
=======
Interactive entry point for ISO Certificate extraction, classification,
and rule-based golden dataset validation with baseline regression tracking.

Modes:
  [1] Run Extraction (Process Excel containing URLs)
  [2] Dataset Validation (Extract golden dataset URLs & compare with baseline)
  [3] Production Pipeline (Extraction + Rule-Based Verification & Scoring)
"""

from __future__ import annotations

import logging
import os
import shutil
import sys
import tempfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from typing import List, Dict, Any, Tuple, Optional

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from extractor import ISOExtractor, RunAnalytics, URLMetrics
from dataset_validation import DatasetValidator
from verifier import CertificateVerifier
from tracing import trace_certificate, span_verification, flush_traces
from database import (
    log_certificate_request,
    DEFAULT_DB_PATH,
    has_reference_data,
    get_reference_records,
    get_recent_production_records,
    get_production_count,
)
from judge import LLMJudge
from drift import execute_drift_test, print_drift_dashboard

# ── Logging Configuration ─────────────────────────────────────────────────────
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%H:%M:%S",
)
logger = logging.getLogger("ISOExtractor")


def prompt_with_default(prompt_text: str, default: str = "") -> str:
    """
    Prompts user for input via input().
    Returns default on empty input or in non-interactive CI environments (EOFError).
    """
    try:
        val = input(prompt_text).strip().strip('"').strip("'")
        return val if val else default
    except (EOFError, OSError):
        # Non-interactive CI execution fallback
        print(f"{prompt_text}{default} (non-interactive default)")
        return default


def load_api_key(key_file: str = "gemini.txt") -> str:
    """
    Reads a single Gemini API key.
    Checks GEMINI_API_KEY environment variable first, then key_file.
    """
    env_key = os.environ.get("GEMINI_API_KEY", "").strip() or os.environ.get("GEMINI_API_KEYS", "").strip()
    if env_key:
        return env_key

    path = Path(key_file)
    if not path.is_absolute():
        candidate = Path.cwd() / key_file
        if not candidate.exists():
            candidate = Path(__file__).resolve().parent / key_file
        path = candidate

    if not path.exists():
        raise FileNotFoundError(
            f"API key file '{key_file}' not found at {path.resolve()}. "
            "Please create it or set the GEMINI_API_KEY environment variable."
        )

    content = path.read_text(encoding="utf-8").strip()
    if not content:
        raise ValueError(f"'{key_file}' is empty. Please add your Gemini API key.")

    for line in content.splitlines():
        line = line.strip()
        if line and not line.startswith("#"):
            return line

    raise ValueError(f"No valid API key found in '{key_file}'.")


def load_input_urls(excel_path: str) -> List[str]:
    """Reads URLs from an Excel sheet. Looks for a column named 'url'."""
    p = Path(excel_path)
    if not p.exists():
        raise FileNotFoundError(f"Input file '{excel_path}' does not exist.")

    df = pd.read_excel(excel_path)

    url_col = None
    for col in df.columns:
        if str(col).strip().lower() == "url":
            url_col = col
            break

    if url_col is None:
        raise ValueError(
            f"Excel file must contain a column called 'url'. Found columns: {list(df.columns)}"
        )

    urls = df[url_col].dropna().astype(str).str.strip().tolist()
    urls = [u for u in urls if u and u.lower() not in ("nan", "none")]
    return urls


def format_excel_output(excel_path: str) -> None:
    """Applies professional styling to the output Excel workbook, including verification badges."""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(excel_path)
        ws = wb.active
        if not ws:
            return

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )
        align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        align_center = Alignment(horizontal="center", vertical="center")

        # Status badge styles
        fill_approved = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        font_approved = Font(name="Calibri", size=11, bold=True, color="276A3C")

        fill_review = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        font_review = Font(name="Calibri", size=11, bold=True, color="8C6B00")

        fill_reject = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
        font_reject = Font(name="Calibri", size=11, bold=True, color="C00000")

        font_score = Font(name="Calibri", size=11, bold=True)

        verif_col_idx = None
        score_col_idx = None

        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_center

            header_text = str(cell.value or "").strip().lower()
            if header_text == "verification_status":
                verif_col_idx = col_idx
            elif header_text == "confidence_score":
                score_col_idx = col_idx

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.border = thin_border
                cell.alignment = align_left

                if verif_col_idx and cell.column == verif_col_idx:
                    val_str = str(cell.value or "").strip()
                    cell.alignment = align_center
                    if val_str == "Approved":
                        cell.fill = fill_approved
                        cell.font = font_approved
                    elif val_str == "Need Review":
                        cell.fill = fill_review
                        cell.font = font_review
                    elif val_str == "Reject":
                        cell.fill = fill_reject
                        cell.font = font_reject

                if score_col_idx and cell.column == score_col_idx:
                    cell.alignment = align_center
                    cell.font = font_score

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                line_len = max((len(l) for l in val.split("\n")), default=0)
                max_len = max(max_len, line_len)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

        wb.save(excel_path)
    except Exception as e:
        logger.warning(f"Could not apply Excel formatting: {e}")


def execute_extraction(
    urls: List[str],
    api_key: str,
    model_name: str = "gemini-2.5-flash-lite",
    workers: int = 4,
    output_path: Optional[str] = None,
    save_analytics: bool = True,
    verifier: Optional[CertificateVerifier] = None,
    enable_tracing: bool = False,
    enable_db_logging: bool = False,
    enable_llm_judge: bool = False,
    judge_sample_ratio: float = 0.15,
    traffic_type: str = "production",
    enable_drift_monitoring: bool = False,
) -> pd.DataFrame:
    """Runs multithreaded extraction on a list of certificate URLs, with optional verifier, tracing, DB logging, LLM judge, and drift detection."""
    extractor = ISOExtractor(api_key=api_key, model_name=model_name)
    run_analytics = RunAnalytics(model_name=model_name)

    temp_dir = tempfile.mkdtemp(prefix="iso_extract_")
    logger.info(f"Using temporary directory: {temp_dir}")

    total_urls = len(urls)
    results: List[Dict[str, Any]] = [None] * total_urls
    metrics_by_url: Dict[str, URLMetrics] = {}
    completed_count = 0
    last_drift_cadence_to_idx: Optional[int] = None

    print(f"\nProcessing {total_urls} certificates with {workers} thread(s) (traffic: {traffic_type})...")

    def process_url_worker(index: int, cert_url: str) -> Tuple[int, Dict[str, Any], URLMetrics]:
        if enable_tracing:
            with trace_certificate(cert_url, metadata={"model": model_name, "index": index + 1}) as cert_trace:
                record, metric = extractor.process_cert_url(cert_url=cert_url, temp_base_dir=temp_dir)
                cert_trace.update(
                    output=record,
                    metadata={
                        "status": record.get("status"),
                        "cost_usd": metric.total_cost_usd,
                        "latency_ms": metric.total_latency_ms,
                    },
                )
                return index, record, metric
        else:
            record, metric = extractor.process_cert_url(cert_url=cert_url, temp_base_dir=temp_dir)
            return index, record, metric

    try:
        with ThreadPoolExecutor(max_workers=workers) as executor:
            future_to_idx = {
                executor.submit(process_url_worker, idx, url): idx
                for idx, url in enumerate(urls)
            }

            for future in as_completed(future_to_idx):
                idx, record, metric = future.result()
                results[idx] = record
                if record.get("url"):
                    metrics_by_url[record["url"]] = metric
                run_analytics.record_url_metrics(metric)
                completed_count += 1

                # 1. Apply single-row verification rule scoring immediately
                if verifier is not None:
                    verif_res = verifier.verify_row(record)
                    record.update(verif_res.to_dict())

                record["traffic_type"] = traffic_type

                # 2. Log to SQLite database immediately
                if enable_db_logging:
                    req_id = log_certificate_request(
                        record=record,
                        metrics=metric,
                        db_path=DEFAULT_DB_PATH,
                    )
                    record["request_id"] = req_id

                status = record.get("status", "Unknown")
                cert_type = record.get("cert_type") or "N/A"
                is_cert = record.get("is_certificate")
                is_iso = record.get("is_iso_iatf")
                cert_lang = record.get("certificate_language") or "N/A"
                lat_s = metric.total_latency_ms / 1000.0
                cost = metric.total_cost_usd

                logger.info(
                    f"[{completed_count}/{total_urls}] Finished URL #{idx+1} | "
                    f"Status: {status} | Lang: {cert_lang} | Latency: {lat_s:.2f}s | Cost: ${cost:.5f} | "
                    f"Cert: {is_cert} | ISO/IATF: {is_iso} | Type: {cert_type}"
                )

                # 3. Intermediate 50-sample drift test cadence
                if enable_drift_monitoring and traffic_type == "production" and (completed_count % 50 == 0):
                    try:
                        ref_records = get_reference_records(DEFAULT_DB_PATH)
                        recent_records = get_recent_production_records(limit=50, db_path=DEFAULT_DB_PATH)
                        if ref_records and len(recent_records) >= 5:
                            total_prod = get_production_count(DEFAULT_DB_PATH)
                            to_idx = total_prod
                            from_idx = max(1, total_prod - len(recent_records) + 1)
                            print(f"\n[DRIFT MONITOR] Running drift test at 50-sample cadence (Samples #{from_idx} to #{to_idx})...")
                            drift_rep, alarms_list = execute_drift_test(
                                ref_records=ref_records,
                                recent_records=recent_records,
                                batch_index=completed_count,
                                from_index=from_idx,
                                to_index=to_idx,
                                save_reports=True,
                                drift_report_path="analytics/drift_test.json",
                                alarms_path="analytics/alarms.json",
                            )
                            last_drift_cadence_to_idx = to_idx
                            print_drift_dashboard(drift_rep, alarms_list)
                    except Exception as drift_err:
                        logger.warning(f"Error during intermediate drift test check: {drift_err}")

    except KeyboardInterrupt:
        logger.warning("Interrupted by user. Collecting partial results...")
    finally:
        try:
            shutil.rmtree(temp_dir, ignore_errors=True)
        except Exception:
            pass

    valid_results = [r for r in results if r is not None]
    if not valid_results:
        logger.warning("No records were successfully processed.")
        return pd.DataFrame()

    df_out = pd.DataFrame(valid_results)

    # ── Optional Verification Summary ──────────────────────────────────────────
    verif_summary = None
    if verifier is not None:
        print("\nFinalizing rule-based verification and confidence scoring...")
        if enable_tracing:
            with span_verification({"batch_size": len(df_out)}) as verif_span:
                df_out, verif_summary = verifier.verify_dataframe(df_out)
                verif_span.update(output=verif_summary)
        else:
            df_out, verif_summary = verifier.verify_dataframe(df_out)

    # ── Post-Batch Stratified LLM-as-a-Judge Evaluation (~15%) ─────────────────
    judge_summary = None
    if enable_llm_judge and len(df_out) > 0:
        records_to_judge = df_out.to_dict(orient="records")
        judge = LLMJudge(api_key=api_key, model_name=model_name)
        judge_summary = judge.evaluate_batch(
            records=records_to_judge,
            sample_ratio=judge_sample_ratio,
            db_path=DEFAULT_DB_PATH,
        )

    # ── Final Post-Batch Drift Check (with updated Judge scores) ───────────────
    if enable_drift_monitoring and traffic_type == "production" and len(df_out) > 0:
        try:
            ref_records = get_reference_records(DEFAULT_DB_PATH)
            recent_records = get_recent_production_records(limit=50, db_path=DEFAULT_DB_PATH)
            if ref_records and recent_records:
                total_prod = get_production_count(DEFAULT_DB_PATH)
                to_idx = total_prod
                from_idx = max(1, total_prod - len(recent_records) + 1)
                if to_idx == last_drift_cadence_to_idx and judge_summary is None:
                    logger.info(f"Skipping redundant post-batch drift check: window #{from_idx}-#{to_idx} already checked.")
                else:
                    print(f"\n[FINAL DRIFT CHECK] Running post-batch drift test with latest production window (Samples #{from_idx} to #{to_idx})...")
                    drift_rep, alarms_list = execute_drift_test(
                        ref_records=ref_records,
                        recent_records=recent_records,
                        batch_index=completed_count,
                        from_index=from_idx,
                        to_index=to_idx,
                        save_reports=True,
                        drift_report_path="analytics/drift_test.json",
                        alarms_path="analytics/alarms.json",
                    )
                    print_drift_dashboard(drift_rep, alarms_list)
        except Exception as drift_err:
            logger.warning(f"Error during post-batch drift test check: {drift_err}")

    # Flush traces to Langfuse if enabled
    if enable_tracing:
        flush_traces()

    ordered_columns = [
        "request_id",
        "traffic_type",
        "url",
        "status",
        "verification_status",
        "confidence_score",
        "verification_reviews",
        "is_certificate",
        "is_iso_iatf",
        "cert_type",
        "screening_reason",
        "certificate_language",
        "image_pixel_size",
        "site_name",
        "address",
        "scope",
        "iso_type",
        "issued_date",
        "expiration_date",
        "site_iso_type",
        "iso_product",
        "site_iso_type_status",
        "doc_type",
        "page_count",
        "latency_ms",
        "cost_usd",
        "total_tokens",
    ]
    existing_cols = [c for c in ordered_columns if c in df_out.columns]
    other_cols = [c for c in df_out.columns if c not in existing_cols]
    df_out = df_out[existing_cols + other_cols]

    if output_path:
        out_p = Path(output_path)
        out_p.parent.mkdir(parents=True, exist_ok=True)
        df_out.to_excel(output_path, index=False)
        format_excel_output(output_path)
        print(f"\nSaved {len(valid_results)} record(s) to: {out_p.resolve()}")

    if save_analytics:
        analytics_file, run_key = run_analytics.save_to_json(
            analytics_dir="analytics",
            filename="analytics.json",
            verification_summary=verif_summary,
            judge_summary=judge_summary,
        )
        summary_data = run_analytics.generate_summary()["summary"]
        print("-" * 65)
        print(f"ANALYTICS SUMMARY [{run_key.upper()}]:")
        print(f"  Total Processed URLs : {summary_data['total_urls']}")
        print(f"  Successful URLs      : {summary_data['successful_urls']}")
        print(f"  Failed / Review URLs : {summary_data['failed_urls']}")
        print(f"  Total Latency        : {summary_data['total_latency_seconds']} s ({summary_data['total_latency_ms']} ms)")
        print(f"  Avg Latency per URL  : {summary_data['avg_latency_per_url_ms']} ms")
        print(f"  Total Tokens         : {summary_data['total_tokens']:,} (Input: {summary_data['total_input_tokens']:,}, Output: {summary_data['total_output_tokens']:,})")
        print(f"  Estimated Cost       : ${summary_data['total_cost_usd']:.5f} USD")
        print(f"  Detailed Analytics   : {analytics_file}")
        print(f"  Aggregated Analytics : analytics/analytics_aggregated.json")

    if verif_summary is not None:
        print("\n" + "=" * 65)
        print("                 VERIFICATION RESULTS SUMMARY")
        print("=" * 65)
        print(f"  Total Cases Evaluated : {verif_summary['total_cases']}")
        print(f"  Approved (Score >= 90): {verif_summary['approved_cases']} ({verif_summary['approval_rate_percent']}%)")
        print(f"  Need Review (50 - 89) : {verif_summary['need_review_cases']}")
        print(f"  Rejected (Score < 50) : {verif_summary['rejected_cases']}")
        print(f"  Average Confidence    : {verif_summary['avg_score']} / 100")
        print("-" * 65)
        if verif_summary.get("rule_violation_counts"):
            print("RULE VIOLATIONS DETECTED:")
            for rule_k, count_v in verif_summary["rule_violation_counts"].items():
                print(f"  • {rule_k:<24}: {count_v} case(s)")
        else:
            print("  All cases fully passed verification rules!")
        print("=" * 65)

    if judge_summary is not None:
        print("\n" + "=" * 65)
        print("                 LLM JUDGE AUDIT SUMMARY")
        print("=" * 65)
        print(f"  Total Sampled Cases   : {judge_summary['total_sampled']} ({judge_summary['sample_rate_percent']}%)")
        print(f"  Passed Cases (Score 1): {judge_summary['passed_count']} ({judge_summary['pass_rate_percent']}%)")
        print(f"  Failed Cases (Score 0): {judge_summary['failed_count']}")
        print(f"  Average Judge Score   : {judge_summary['avg_judge_score']:.4f} / 1.00")
        print("-" * 65)
        print("STRATIFIED VERDICT BREAKDOWN:")
        for cat_name, cat_stats in judge_summary.get("category_breakdown", {}).items():
            print(f"  • {cat_name:<16}: {cat_stats['passed']}/{cat_stats['sampled']} passed ({cat_stats['pass_rate_percent']}%)")
        print("=" * 65)

    return df_out


def run_mode_extraction():
    """Mode 1: Standard extraction from an Excel sheet containing URLs."""
    print("\n--- Mode 1: Standard Extraction ---")
    input_path = prompt_with_default("Enter path to input Excel file: ")
    if not input_path:
        print("Error: Input file path cannot be empty.")
        sys.exit(1)

    output_path = prompt_with_default(
        "Enter path for output Excel file (default: output.xlsx): ",
        default="output.xlsx",
    )
    if not output_path.lower().endswith(".xlsx"):
        output_path += ".xlsx"

    model_name = prompt_with_default(
        "Enter Gemini model [default: gemini-2.5-flash-lite]: ",
        default="gemini-2.5-flash-lite",
    )
    threads_input = prompt_with_default(
        "Enter number of worker threads [default: 4]: ",
        default="4",
    )
    try:
        workers = int(threads_input) if threads_input else 4
        if workers < 1:
            workers = 1
    except ValueError:
        workers = 4

    try:
        api_key = load_api_key("gemini.txt")
    except Exception as e:
        logger.error(f"Failed to load API key: {e}")
        sys.exit(1)

    try:
        urls = load_input_urls(input_path)
        logger.info(f"Loaded {len(urls)} URL(s) from '{input_path}'")
    except Exception as e:
        logger.error(f"Failed to load input URLs: {e}")
        sys.exit(1)

    execute_extraction(
        urls=urls,
        api_key=api_key,
        model_name=model_name,
        workers=workers,
        output_path=output_path,
        save_analytics=True,
    )


def run_mode_validation():
    """Mode 2: Rule-based validation running extraction on the golden dataset and comparing with baseline."""
    print("\n--- Mode 2: Dataset Validation & Baseline Tracking ---")
    golden_path = prompt_with_default(
        "Enter path to golden dataset (default: golden_dataset_regression.xlsx): ",
        default="golden_dataset_regression.xlsx",
    )

    if not Path(golden_path).exists():
        print(f"Error: Golden dataset '{golden_path}' not found.")
        sys.exit(1)

    model_name = prompt_with_default(
        "Enter Gemini model [default: gemini-2.5-flash-lite]: ",
        default="gemini-2.5-flash-lite",
    )
    threads_input = prompt_with_default(
        "Enter number of worker threads [default: 4]: ",
        default="4",
    )
    try:
        workers = int(threads_input) if threads_input else 4
        if workers < 1:
            workers = 1
    except ValueError:
        workers = 4

    try:
        api_key = load_api_key("gemini.txt")
    except Exception as e:
        logger.error(f"Failed to load API key: {e}")
        sys.exit(1)

    # Directly run extraction on the golden dataset URLs with Verifier enabled
    urls = load_input_urls(golden_path)
    raw_output_path = "dataset_validation/golden_extracted_output.xlsx"
    verifier = CertificateVerifier(initial_score=100)
    predicted_df = execute_extraction(
        urls=urls,
        api_key=api_key,
        model_name=model_name,
        workers=workers,
        output_path=raw_output_path,
        save_analytics=True,
        verifier=verifier,
    )

    if predicted_df is None or predicted_df.empty:
        print("Error: Extraction produced no records to validate.")
        sys.exit(1)

    print("\nRunning rule-based dataset validation against golden dataset...")
    validator = DatasetValidator(golden_path=golden_path)
    val_results = validator.validate(predicted_df=predicted_df)

    # ── Compare with Baseline (analytics/baseline.json) ────────────────────────
    baseline_cmp = validator.compare_with_baseline(
        current_summary=val_results["summary"],
        analytics_dir="analytics",
        baseline_filename="baseline.json",
        new_baseline_filename="new_baseline.json",
    )

    report_excel = prompt_with_default(
        "Enter output path for validation Excel report (default: dataset_validation/validation_report.xlsx): ",
        default="dataset_validation/validation_report.xlsx",
    )
    if not report_excel.lower().endswith(".xlsx"):
        report_excel += ".xlsx"

    report_json = "dataset_validation/validation_summary.json"
    saved_excel, saved_json = validator.save_reports(val_results, report_excel, report_json)

    summary = val_results["summary"]
    print("\n" + "=" * 70)
    print("                 DATASET VALIDATION RESULTS")
    print("=" * 70)
    print(f"Golden File           : {summary['golden_file']}")
    print(f"Total Cases Evaluated : {summary['total_cases']}")
    print(f"Passed Cases          : {summary['passed_cases']} / {summary['total_cases']}")
    print(f"Failed Cases          : {summary['failed_cases']}")
    print(f"Pass Rate (%)         : {summary['pass_rate_percent']}%")
    print(f"Avg Total Score       : {summary['average_total_score']} / 10 pts")
    print(f"Avg Row Score         : {summary['average_row_score']} (scale: 0.0 - 1.0)")
    print("-" * 70)
    print("FALSE PASS METRICS (Verifier Approved vs Validator Failed):")
    print(f"  • False Pass Cases    : {summary.get('false_pass_cases', 0)} / {summary['total_cases']} ({summary.get('false_pass_rate_total_percent', 0.0)}% of total)")
    print(f"  • Rate of Failed Cases: {summary.get('false_pass_rate_failed_percent', 0.0)}% of failed cases")
    if summary.get("false_pass_urls"):
        print("  • False Pass URLs:")
        for fp in summary["false_pass_urls"]:
            print(f"    - Row #{fp['index']}: {fp['url']} (Confidence: {fp['confidence_score']})")
    print("-" * 70)
    print("FEATURE SCORES (AVERAGE PER COLUMN):")
    for feat_name, score_val in summary["feature_scores"].items():
        pct = round(score_val * 100, 2)
        print(f"  • {feat_name:28}: {score_val:.4f}  ({pct:5.1f}% accuracy)")

    print("-" * 70)
    # ── Baseline Comparison Output ────────────────────────────────────────────
    if not baseline_cmp["has_baseline"]:
        print(f"[BASELINE CREATED] Initial baseline saved to:")
        print(f"  --> {baseline_cmp['baseline_path']}")
    else:
        print("BASELINE REGRESSION COMPARISON:")
        print(f"Reference Baseline: {baseline_cmp['baseline_path']}")
        print(f"{'Feature':<28} {'Baseline':>10} {'New Run':>10} {'Delta':>8} {'Status'}")
        print("-" * 70)
        for feat, data in baseline_cmp["feature_comparisons"].items():
            print(
                f"{feat:<28} {data['baseline']:>10.4f} {data['new']:>10.4f} "
                f"{data['delta']:>+8d} {data['status']}"
            )
        print("-" * 70)
        total_delta = baseline_cmp["total_delta"]
        sign = "+" if total_delta > 0 else ""
        print(f"TOTAL FEATURE IMPROVEMENT DELTA: {sign}{total_delta}")

        if baseline_cmp["saved_new_baseline"]:
            print("\n[IMPROVEMENT DETECTED] Overall score is positive!")
            print(f"Saved candidate baseline to:")
            print(f"  --> {baseline_cmp['new_baseline_path']}")
            print("(Note: The active baseline remains baseline.json; new_baseline.json is saved for your review.)")
        else:
            print(f"\n[NO REGRESSION IMPROVEMENT] Total Delta is {sign}{total_delta} (not positive).")
            print("new_baseline.json was not generated.")

    print("-" * 70)
    print("Reports Generated:")
    print(f"  Excel Report : {saved_excel}")
    print(f"  JSON Summary : {saved_json}")
    print("=" * 70)


def run_mode_production():
    """Mode 3: Production Pipeline (Extraction + Rule-Based Verification & Scoring)."""
    print("\n--- Mode 3: Production Pipeline (Extraction + Verification & Scoring) ---")
    input_path = prompt_with_default("Enter path to input Excel file: ")
    if not input_path:
        print("Error: Input file path cannot be empty.")
        sys.exit(1)

    output_path = prompt_with_default(
        "Enter path for verified output Excel file (default: production_output.xlsx): ",
        default="production_output.xlsx",
    )
    if not output_path.lower().endswith(".xlsx"):
        output_path += ".xlsx"

    model_name = prompt_with_default(
        "Enter Gemini model [default: gemini-2.5-flash-lite]: ",
        default="gemini-2.5-flash-lite",
    )
    threads_input = prompt_with_default(
        "Enter number of worker threads [default: 4]: ",
        default="4",
    )
    try:
        workers = int(threads_input) if threads_input else 4
        if workers < 1:
            workers = 1
    except ValueError:
        workers = 4

    try:
        api_key = load_api_key("gemini.txt")
    except Exception as e:
        logger.error(f"Failed to load API key: {e}")
        sys.exit(1)

    try:
        urls = load_input_urls(input_path)
        logger.info(f"Loaded {len(urls)} URL(s) from '{input_path}'")
    except Exception as e:
        logger.error(f"Failed to load input URLs: {e}")
        sys.exit(1)

    verifier = CertificateVerifier(initial_score=100)

    # ── 1. Check & Ingest Reference Set if not already present in SQLite ───────
    if not has_reference_data(DEFAULT_DB_PATH):
        ref_file = "reference_set.xlsx"
        if os.path.exists(ref_file):
            print("\n" + "=" * 75)
            print(f"[DRIFT BASELINE] No reference baseline found in SQLite ({DEFAULT_DB_PATH}).")
            print(f"Ingesting baseline certificates from '{ref_file}' before production batch...")
            print("=" * 75)
            try:
                ref_urls = load_input_urls(ref_file)
                execute_extraction(
                    urls=ref_urls,
                    api_key=api_key,
                    model_name=model_name,
                    workers=workers,
                    output_path=None,
                    save_analytics=False,
                    verifier=verifier,
                    enable_tracing=True,
                    enable_db_logging=True,
                    enable_llm_judge=True,
                    judge_sample_ratio=0.15,
                    traffic_type="reference",
                    enable_drift_monitoring=False,
                )
                print("\n[DRIFT BASELINE] Reference set baseline successfully ingested into SQLite.")
            except Exception as ref_err:
                logger.warning(f"[DRIFT BASELINE] Could not ingest reference set: {ref_err}")
        else:
            logger.warning(f"[DRIFT BASELINE] '{ref_file}' not found. Intermediate drift checks will use available historical data.")
    else:
        logger.info(f"[DRIFT BASELINE] Reference set data already present in SQLite ({DEFAULT_DB_PATH}).")

    # ── 2. Run Production Extraction & Drift Monitoring ───────────────────────
    execute_extraction(
        urls=urls,
        api_key=api_key,
        model_name=model_name,
        workers=workers,
        output_path=output_path,
        save_analytics=True,
        verifier=verifier,
        enable_tracing=True,
        enable_db_logging=True,
        enable_llm_judge=True,
        judge_sample_ratio=0.15,
        traffic_type="production",
        enable_drift_monitoring=True,
    )


def main():
    print("=" * 65)
    print("      ISO / IATF Certificate Extraction & Validation")
    print("=" * 65)

    mode_choice = None
    valid_modes = ("1", "2", "3", "production", "prod")
    if len(sys.argv) > 1 and sys.argv[1].strip() in valid_modes:
        mode_choice = sys.argv[1].strip()
    elif os.environ.get("RUN_MODE") in valid_modes:
        mode_choice = os.environ.get("RUN_MODE")
    else:
        print("Select Mode:")
        print("  [1] Run Extraction (Process Excel containing URLs)")
        print("  [2] Dataset Validation (Extract golden dataset & compare with baseline)")
        print("  [3] Production Pipeline (Extraction + Rule-Based Verification & Scoring)")
        mode_choice = prompt_with_default("Choose mode [1, 2, or 3, default: 3]: ", default="3")

    if mode_choice == "2":
        run_mode_validation()
    elif mode_choice in ("3", "production", "prod"):
        run_mode_production()
    else:
        run_mode_extraction()


if __name__ == "__main__":
    main()

