#!/usr/bin/env python3
"""
export_db_to_excel.py
=====================
Exports all certificate request logs from the SQLite database
(analytics/certificate_logs.db) into a styled Excel file.
"""

from __future__ import annotations

import os
from pathlib import Path
import sqlite3
import sys

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DEFAULT_DB_PATH = "analytics/certificate_logs.db"
DEFAULT_OUTPUT_PATH = "analytics/certificate_logs_export.xlsx"


def format_excel(excel_path: str):
    """Applies clean styling and auto-fits column widths."""
    from openpyxl import load_workbook

    wb = load_workbook(excel_path)
    ws = wb.active

    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    fill_approved = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")  # soft green
    font_approved = Font(color="375623", bold=True)
    fill_review = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")    # soft yellow
    font_review = Font(color="7F6000", bold=True)
    fill_reject = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")    # soft red
    font_reject = Font(color="C65911", bold=True)

    verdict_col_idx = None
    judge_col_idx = None

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if cell.value == "verdict":
            verdict_col_idx = cell.column
        elif cell.value == "judge_score":
            judge_col_idx = cell.column

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = thin_border
            if verdict_col_idx and cell.column == verdict_col_idx:
                val = str(cell.value or "")
                if val == "Approved":
                    cell.fill = fill_approved
                    cell.font = font_approved
                elif val == "Need Review":
                    cell.fill = fill_review
                    cell.font = font_review
                elif val == "Reject":
                    cell.fill = fill_reject
                    cell.font = font_reject

            if judge_col_idx and cell.column == judge_col_idx and cell.value is not None:
                try:
                    score = int(cell.value)
                    if score == 1:
                        cell.fill = fill_approved
                        cell.font = font_approved
                    elif score == 0:
                        cell.fill = fill_reject
                        cell.font = font_reject
                except (ValueError, TypeError):
                    pass

    # Auto-adjust column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            line_len = max((len(line) for line in val_str.split("\n")), default=0)
            max_len = max(max_len, line_len)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    wb.save(excel_path)


def export_db_to_excel(
    db_path: str = DEFAULT_DB_PATH,
    output_excel_path: str = DEFAULT_OUTPUT_PATH,
) -> str:
    """Reads all rows from SQLite database and writes them to an Excel spreadsheet."""
    db_file = Path(db_path)
    if not db_file.exists():
        print(f"[ERROR] Database file not found at: {db_file.resolve()}")
        print("Please run Mode 3 first to generate certificate logs in the database.")
        sys.exit(1)

    conn = sqlite3.connect(str(db_file))
    try:
        query = "SELECT * FROM certificate_requests ORDER BY id ASC;"
        df = pd.read_sql_query(query, conn)
    finally:
        conn.close()

    total_rows = len(df)
    if total_rows == 0:
        print(f"[INFO] Database at '{db_path}' is empty (0 records found).")
        return ""

    out_file = Path(output_excel_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)

    df.to_excel(str(out_file), index=False)
    try:
        format_excel(str(out_file))
    except Exception as e:
        print(f"[WARNING] Could not apply Excel formatting: {e}")

    print("\n" + "=" * 65)
    print("           SQLITE DATABASE EXPORT COMPLETED")
    print("=" * 65)
    print(f"Source Database   : {db_file.resolve()}")
    print(f"Total Rows Found  : {total_rows}")
    if "traffic_type" in df.columns:
        ref_n = int((df["traffic_type"] == "reference").sum())
        prod_n = int((df["traffic_type"] == "production").sum())
        print(f"Traffic Breakdown : Reference: {ref_n} | Production: {prod_n}")
    judged_count = int(df["judge_score"].notna().sum()) if "judge_score" in df.columns else 0
    print(f"Judged Rows       : {judged_count} / {total_rows}")
    print(f"Saved Excel File  : {out_file.resolve()}")
    print("=" * 65)

    return str(out_file.resolve())


def main():
    db_path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_DB_PATH
    output_path = sys.argv[2] if len(sys.argv) > 2 else DEFAULT_OUTPUT_PATH
    export_db_to_excel(db_path, output_path)


if __name__ == "__main__":
    main()
